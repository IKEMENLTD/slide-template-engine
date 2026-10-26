"""Phase 5：手順書（docs/TEMPLATE_GUIDE.md）11章の例「引用」の型を、手順どおりに足してみる。
PowerPoint の操作の代わりに python-pptx で、手順書に書いてあることだけを行う（写真＝画像を挿入して図形でトリミング、
文字＝テキストボックス・自動調整なし・行間固定値、飾り＝アクセント2の平行四辺形、ノート＝YAML）。
使い方: python3 scripts/phase5/guide-dryrun.py → out/phase5/guide-dryrun/provisional+quote.pptx
"""
import importlib.util
import shutil
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("tpl", ROOT / "scripts/build-provisional-template.py")
tpl = importlib.util.module_from_spec(spec)
spec.loader.exec_module(tpl)

OUT = ROOT / "out/phase5/guide-dryrun"
OUT.mkdir(parents=True, exist_ok=True)
dst = OUT / "provisional+quote.pptx"
shutil.copy(ROOT / "templates/provisional.pptx", dst)
prs = Presentation(str(dst))
# 1. 白紙のスライドを足す
s = prs.slides.add_slide(prs.slide_layouts[6])
ids = tpl.Ids(s)
# 飾り：アクセント2 の平行四辺形（名前は付けない＝飾り）
tpl.poly(s, ids, "平行四辺形 1", [(430, 400), (500, 330), (900, 330), (830, 400)], "accent2")
# 2. 写真：画像を挿入 → 図形（平行四辺形）でトリミング → 名前 img.photo
tpl.picture(s, ids, "img.photo", tpl.ASSETS / "side.jpg", 0, 0, 420, 540, pts=[(0, 0), (420, 0), (300, 540), (0, 540)], descr="引用の写真")
# 5. 見出し
tpl.text(s, ids, "txt.headline", 460, 50, 460, 56, "ダミー見出し", ea=tpl.MINCHO, latin=tpl.MINCHO, size=32.5, line=40, anchor="ctr")
# 3. 引用文：明朝・28pt・行間 固定値 40pt・自動調整なし
tpl.text(s, ids, "txt.quote", 460, 140, 440, 170, "ダミーの引用文", ea=tpl.MINCHO, latin=tpl.MINCHO, size=28, line=40)
# 4. 肩書き：ゴシック・14pt・行間 固定値 20pt・自動調整なし
tpl.text(s, ids, "txt.source", 520, 345, 300, 44, "ダミー肩書き", ea=tpl.GOTHIC, latin=tpl.LATIN, size=14, line=20)
# 6. ノート
tpl.notes(s, "type: quote\nrole: [emphasis]\nlabel: 引用（お客様の声）")
prs.save(str(dst))
print(dst)
