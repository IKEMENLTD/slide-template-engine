"""Phase 0 検証④：文字幅の予測と、実際の折り返しを比べるための PPTX を作る。

1スライドに複数の文字枠を置く。枠ごとに幅・フォント・大きさ・文言を変える。
各枠の条件は cases.json に書き出し、fontkit 側（measure-wrap.ts）が同じ条件で予測する。

使い方: python3 scripts/phase0/build_wrap_probe.py <出力.pptx> <cases.json>
"""

import json
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"

TEXTS = [
    "業務が属人化しており、自分がいないと回らない。採用コストも教育コストも高騰している。",
    "見積もりを取ったら数百万〜数千万と言われた。補助金は書類が多く、何が自社に使えるか分からない。",
    "省力化投資補助金は2024年に新設された制度で、最大補助額は2,000万円、補助率は2/3です。",
    "「丸投げOK」で申請代行の費用はゼロ。不採択時は0円です（2025年度対応版）。",
    "社長が現場から離れられないという悩みは、業務システムの導入で解決できます。",
    "DX推進とAI活用でBackOffice業務を効率化し、Excel管理から脱却する。",
]
FONTS = [
    ("Noto Sans JP", False, "fonts/NotoSansJP-Regular.otf"),
    ("Noto Sans JP", True, "fonts/NotoSansJP-Bold.otf"),
    ("Noto Sans JP Medium", False, "fonts/NotoSansJP-Medium.otf"),
    ("Noto Serif JP", False, "fonts/NotoSerifJP-Regular.otf"),
]
SIZES = [12, 16, 20]
WIDTHS = [180, 260, 340]  # pt


def main(out: str, cases_out: str) -> None:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    cases = []
    slide = None
    y = 0.0
    n = 0
    for fi, (face, bold, file) in enumerate(FONTS):
        for si, size in enumerate(SIZES):
            for wi, width in enumerate(WIDTHS):
                text = TEXTS[(fi + si + wi) % len(TEXTS)]
                height = size * 1.6 * 6
                if slide is None or y + height > 520:
                    slide = prs.slides.add_slide(prs.slide_layouts[6])
                    y = 10.0
                    col = 0
                x = 10
                tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
                name = f"case.{len(cases)}"
                tb.name = name
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                bp = tf._txBody.find(f"{{{A}}}bodyPr")
                etree.SubElement(bp, f"{{{A}}}noAutofit")
                r = tf.paragraphs[0].add_run()
                r.text = text
                r.font.size = Pt(size)
                r.font.bold = bold
                rpr = r._r.get_or_add_rPr()
                rpr.set("lang", "ja-JP")
                etree.SubElement(rpr, f"{{{A}}}latin", typeface=face)
                etree.SubElement(rpr, f"{{{A}}}ea", typeface=face)
                cases.append({
                    "name": name, "slide": len(prs.slides), "x": x, "y": y, "width": width,
                    "height": height, "size": size, "face": face, "bold": bold, "fontFile": file, "text": text,
                })
                n += 1
                y += height + 10
    prs.save(out)
    with open(cases_out, "w", encoding="utf-8") as f:
        json.dump(cases, f, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
