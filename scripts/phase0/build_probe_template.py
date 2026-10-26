"""Phase 0 用の検証テンプレートを作る。

本番の仮テンプレート（Phase 1）ではない。Phase 0 で確かめたい4点
（見本スライドの複製・日本語の差し替えで書式が残るか・平行四辺形に
切り抜いた写真の差し替え・テーマ色の参照）を、最小の図形で試すためのもの。

色はすべてテーマ色（schemeClr）で参照し、srgbClr を直書きしない。
テーマの値は 4-2 の確定値（青 #0064D6・黄 #E4E74D・灰 #ECECEC・濃灰 #3D3D3D）。

使い方: python3 scripts/phase0/build_probe_template.py <出力.pptx> <ダミー写真.png>
"""

import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
}

# 先方 PDF は 1920x1080pt。PowerPoint の標準ワイド（13.333in x 7.5in = 960x540pt）に
# 0.5 倍で写す。座標は PDF の pt を半分にした値で書く。
SLIDE_W = Emu(12192000)
SLIDE_H = Emu(6858000)


def pt(v: float) -> Emu:
    return Emu(int(v * 12700))


THEME = {
    "dk1": "000000",
    "lt1": "FFFFFF",
    "dk2": "3D3D3D",
    "lt2": "ECECEC",
    "accent1": "0064D6",
    "accent2": "E4E74D",
}


def set_theme(prs: Presentation) -> None:
    master = prs.slide_masters[0]
    theme_part = None
    for rel in master.part.rels.values():
        if rel.reltype.endswith("/theme"):
            theme_part = rel.target_part
    root = etree.fromstring(theme_part.blob)
    clr = root.find(".//a:clrScheme", NS)
    clr.set("name", "ProbeBlueYellow")
    for name, hexv in THEME.items():
        el = clr.find(f"a:{name}", NS)
        for c in list(el):
            el.remove(c)
        etree.SubElement(el, f"{{{NS['a']}}}srgbClr", val=hexv)
    fonts = root.find(".//a:fontScheme", NS)
    fonts.set("name", "ProbeFonts")
    for kind in ("majorFont", "minorFont"):
        f = fonts.find(f"a:{kind}", NS)
        f.find("a:latin", NS).set("typeface", "Barlow")
        f.find("a:ea", NS).set("typeface", "Noto Sans JP")
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def scheme_fill(shape, name: str) -> None:
    """図形の塗りをテーマ色参照にする（srgbClr を書かない）。"""
    sppr = shape._element.spPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        for el in sppr.findall(tag, NS):
            sppr.remove(el)
    fill = etree.SubElement(sppr, f"{{{NS['a']}}}solidFill")
    etree.SubElement(fill, f"{{{NS['a']}}}schemeClr", val=name)
    # 塗りは prstGeom の直後に置く必要がある
    geom = sppr.find("a:prstGeom", NS)
    sppr.remove(fill)
    geom.addnext(fill)
    ln = sppr.find("a:ln", NS)
    if ln is None:
        ln = etree.SubElement(sppr, f"{{{NS['a']}}}ln")
    for c in list(ln):
        ln.remove(c)
    etree.SubElement(ln, f"{{{NS['a']}}}noFill")


def text_box(slide, name, x, y, w, h, text, *, font, size, bold=False, color="tx1", anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    # 余白を 0 にして、枠の幅＝文字が並べる幅にする（測定と一致させるため）
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    bodypr = tf._txBody.find("a:bodyPr", NS)
    # 自動調整（文字を縮める）を明示的に切る。原則4「字を小さくしない」
    for c in list(bodypr):
        bodypr.remove(c)
    etree.SubElement(bodypr, f"{{{NS['a']}}}noAutofit")
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    rpr = r._r.get_or_add_rPr()
    rpr.set("lang", "ja-JP")
    sf = etree.SubElement(rpr, f"{{{NS['a']}}}solidFill")
    etree.SubElement(sf, f"{{{NS['a']}}}schemeClr", val=color)
    etree.SubElement(rpr, f"{{{NS['a']}}}latin", typeface=font[0])
    etree.SubElement(rpr, f"{{{NS['a']}}}ea", typeface=font[1])
    return tb


def set_notes(slide, yaml_text: str) -> None:
    slide.notes_slide.notes_text_frame.text = yaml_text


def picture_parallelogram(slide, name, img, x, y, w, h, adj=50000):
    pic = slide.shapes.add_picture(img, pt(x), pt(y), pt(w), pt(h))
    pic.name = name
    geom = pic._element.spPr.find("a:prstGeom", NS)
    geom.set("prst", "parallelogram")
    av = geom.find("a:avLst", NS)
    if av is None:
        av = etree.SubElement(geom, f"{{{NS['a']}}}avLst")
    gd = etree.SubElement(av, f"{{{NS['a']}}}gd", name="adj", fmla=f"val {adj}")
    return pic


def main(out: str, dummy_img: str) -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    set_theme(prs)
    blank = prs.slide_layouts[6]

    # --- スライド1: 課題提示（2件）を模した型 ---
    s = prs.slides.add_slide(blank)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, pt(368), pt(14), pt(276), pt(28))
    band.name = "deco.title-band"
    scheme_fill(band, "bg2")  # lt2 = #ECECEC
    band.left, band.top, band.width, band.height = pt(368), pt(14), pt(276 * 2), pt(28)
    tri = s.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, pt(900), pt(480), pt(60), pt(60))
    tri.name = "deco.corner-triangle"
    tri.rotation = 270
    scheme_fill(tri, "accent2")
    para = s.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, pt(0), pt(0), pt(170), pt(70))
    para.name = "deco.eyebrow-band"
    scheme_fill(para, "accent1")
    text_box(s, "txt.headline", 386, 14, 520, 28, "ダミー見出しダミー見出し", font=("Noto Serif JP", "Noto Serif JP"), size=20)
    picture_parallelogram(s, "img.hero", dummy_img, 40, 120, 300, 300)
    text_box(s, "txt.item.1", 380, 120, 220, 44, "ダミー項目一", font=("Noto Sans JP", "Noto Sans JP"), size=16, bold=True)
    text_box(s, "txt.item.2", 640, 120, 220, 44, "ダミー項目二", font=("Noto Sans JP", "Noto Sans JP"), size=16, bold=True)
    text_box(s, "txt.body", 380, 200, 480, 120, "ダミー本文。", font=("Noto Sans JP", "Noto Sans JP"), size=12)
    text_box(s, "num.rate", 380, 340, 220, 80, "00", font=("Barlow", "Noto Sans JP"), size=60, color="tx1")
    set_notes(s, 'type: probe-issues\nvariant: "2"\nrole: [problem]\nlabel: 検証用（課題2件）')

    # --- スライド2: 中扉を模した型 ---
    s2 = prs.slides.add_slide(blank)
    big = s2.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, pt(340), pt(120), pt(560), pt(72))
    big.name = "deco.section-band"
    scheme_fill(big, "accent1")
    t = text_box(s2, "txt.headline", 380, 130, 480, 52, "ダミー中扉", font=("Noto Sans JP", "Noto Sans JP"), size=28, bold=True, color="bg1", anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    set_notes(s2, "type: probe-section\nrole: [section]\nlabel: 検証用（中扉）")

    prs.save(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
