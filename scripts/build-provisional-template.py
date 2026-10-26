"""仮テンプレート（templates/provisional.pptx）を作る。Phase 1・Phase 5（D-018）。

先方 PDF（docs/source/0925_資料.pdf）の各ページを、PowerPoint 標準ワイド（960x540pt）に
0.5 倍で写す（D-002）。座標は PDF から抜き出した値の半分。
  p1 表紙／p2 課題提示／p3 数値訴求／p5 注意・強調／p6 中扉／p7 画像が並ぶ（手順）／
  p8 写真コラージュ（結び）／p9 目次／p10 写真＋本文A／p11 写真＋本文B（p4 は説明のページなので型にしない）
p5〜p11 は文言の欄が見本に無い（または見出しだけ）なので、欄の位置は空いている面に当方が置いた（D-018）。

守っていること（REQUIREMENTS 6章・DECISIONS）
- 1 スライド＝1 つの型。ノート欄の先頭に型の情報（YAML）を書く
- 文言を入れる図形は txt.* / num.*、写真は img.*、それ以外は飾り（エンジンは触らない）
- 色はテーマ色の参照（schemeClr）だけ。srgbClr を直書きしない
- 行間は固定値（pt）で持つ。行数の計算を決まった値にするため（D-013）
- 自動調整（文字を縮める）は切る（noAutofit）。字を小さくして収めない
- 文字の外側の余白は 0。枠の幅＝文字が並べる幅

この .py はテンプレートの「作り方」なので、色・座標・字体を書いてよい唯一の場所。
エンジン（src/）にはこれらを書かない。

使い方: python3 scripts/build-provisional-template.py [出力.pptx]
"""

import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parent.parent
ASSETS = ROOT / "templates/provisional-assets"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = {"a": A, "p": P, "r": R}

# ---- テーマ（4-2 の確定値。PDF を ICC で描いた色と一致） ----
THEME = {"dk1": "000000", "lt1": "FFFFFF", "dk2": "3D3D3D", "lt2": "ECECEC", "accent1": "0064D6", "accent2": "E4E74D"}
# 濃淡（4-5「同系色の濃淡」）。今回の3ページでは使わないが、型を足すときのために置く
THEME_EXTRA = {"accent3": "8FBDF0", "accent4": "003F8A", "accent5": "F7F8D0", "accent6": "9E9E9E"}

# ---- 字体（D-005・D-008） ----
MINCHO = "Noto Serif JP"          # 大見出し（先方はリュウミン L）
GOTHIC = "Noto Sans JP"           # 太字はこれに b=1（先方は Noto Sans CJK Bold / 太ゴB101）
GOTHIC_M = "Noto Sans JP Medium"  # Medium（先方は Noto Sans CJK Medium）
LATIN = "Barlow"                  # DIN Next LT Pro の代わり（OFL）
LATIN_M = "Barlow Medium"


def E(pt: float) -> int:
    return int(round(pt * 12700))


def q(tag: str) -> str:
    pre, local = tag.split(":")
    return f"{{{NSMAP[pre]}}}{local}"


def sub(parent, tag, **attrs):
    el = etree.SubElement(parent, q(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


# ---------------------------------------------------------------- テーマ
def set_theme(prs):
    master = prs.slide_masters[0]
    part = next(r.target_part for r in master.part.rels.values() if r.reltype.endswith("/theme"))
    root = etree.fromstring(part.blob)
    clr = root.find(f".//{q('a:clrScheme')}")
    clr.set("name", "Provisional Blue Yellow")
    for name, hexv in {**THEME, **THEME_EXTRA}.items():
        el = clr.find(q(f"a:{name}"))
        for c in list(el):
            el.remove(c)
        sub(el, "a:srgbClr", val=hexv)
    fonts = root.find(f".//{q('a:fontScheme')}")
    fonts.set("name", "Provisional")
    for kind in ("majorFont", "minorFont"):
        f = fonts.find(q(f"a:{kind}"))
        f.find(q("a:latin")).set("typeface", LATIN)
        f.find(q("a:ea")).set("typeface", MINCHO if kind == "majorFont" else GOTHIC)
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------- 図形
class Ids:
    """図形IDはスライドの中で一意にする（python-pptx の add_picture も ID を振るので、毎回いまの最大値から取る）。"""

    def __init__(self, slide=None):
        self.slide = slide

    def next(self):
        ids = [int(e.get("id")) for e in self.slide.shapes._spTree.iter(q("p:cNvPr"))]
        return max(ids + [1]) + 1


def _solid(parent, scheme, alpha=None):
    f = sub(parent, "a:solidFill")
    c = sub(f, "a:schemeClr", val=scheme)
    if alpha is not None:
        sub(c, "a:alpha", val=int(alpha * 1000))
    return f


def _custgeom(sppr, pts, x0, y0, w, h):
    geom = sub(sppr, "a:custGeom")
    for t in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        sub(geom, t)
    sub(geom, "a:rect", l="l", t="t", r="r", b="b")
    pl = sub(geom, "a:pathLst")
    path = sub(pl, "a:path", w=E(w), h=E(h))
    for i, (x, y) in enumerate(pts):
        el = sub(path, "a:moveTo" if i == 0 else "a:lnTo")
        sub(el, "a:pt", x=E(x - x0), y=E(y - y0))
    sub(path, "a:close")


def poly(slide, ids, name, pts, scheme, alpha=None):
    """飾りの多角形。pts は 960x540pt の座標（PDF の座標の半分）。"""
    xs, ys = [p[0] for p in pts], [p[1] for p in pts]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    tree = slide.shapes._spTree
    sp = sub(tree, "p:sp")
    nv = sub(sp, "p:nvSpPr")
    sub(nv, "p:cNvPr", id=ids.next(), name=name)
    sub(nv, "p:cNvSpPr")
    sub(nv, "p:nvPr")
    sppr = sub(sp, "p:spPr")
    xf = sub(sppr, "a:xfrm")
    sub(xf, "a:off", x=E(x0), y=E(y0))
    sub(xf, "a:ext", cx=E(x1 - x0), cy=E(y1 - y0))
    _custgeom(sppr, pts, x0, y0, x1 - x0, y1 - y0)
    _solid(sppr, scheme, alpha)
    ln = sub(sppr, "a:ln")
    sub(ln, "a:noFill")
    return sp


def rect(slide, ids, name, x, y, w, h, scheme, alpha=None):
    return poly(slide, ids, name, [(x, y), (x + w, y), (x + w, y + h), (x, y + h)], scheme, alpha)


def _custgeom_multi(sppr, shapes, x0, y0, w, h):
    """いくつかの多角形を1つの形にする（写真コラージュの菱形）。"""
    geom = sub(sppr, "a:custGeom")
    for t in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        sub(geom, t)
    sub(geom, "a:rect", l="l", t="t", r="r", b="b")
    pl = sub(geom, "a:pathLst")
    for pts in shapes:
        path = sub(pl, "a:path", w=E(w), h=E(h))
        for i, (x, y) in enumerate(pts):
            el = sub(path, "a:moveTo" if i == 0 else "a:lnTo")
            sub(el, "a:pt", x=E(x - x0), y=E(y - y0))
        sub(path, "a:close")


def picture(slide, ids, name, img, x, y, w, h, pts=None, descr="", multi=None):
    """写真枠。pts を渡すと、その多角形で写真を切り抜く（custGeom）。multi は多角形の並び。"""
    pic = slide.shapes.add_picture(str(img), Emu(E(x)), Emu(E(y)), Emu(E(w)), Emu(E(h)))
    pic.name = name
    if descr:
        pic._element.find(f".//{q('p:cNvPr')}").set("descr", descr)
    if pts or multi:
        sppr = pic._element.spPr
        geom = sppr.find(q("a:prstGeom"))
        idx = list(sppr).index(geom)
        sppr.remove(geom)
        tmp = etree.Element(q("a:tmp"))
        if multi:
            _custgeom_multi(tmp, multi, x, y, w, h)
        else:
            _custgeom(tmp, pts, x, y, w, h)
        sppr.insert(idx, tmp[0])
    # 縦横比を変えずに枠を埋める（中央で切り抜く）
    from PIL import Image
    iw, ih = Image.open(img).size
    fr, ir = w / h, iw / ih
    blip = pic._element.find(q("p:blipFill"))
    src = blip.find(q("a:srcRect"))
    if src is None:
        src = etree.Element(q("a:srcRect"))
        blip.insert(1, src)
    if ir > fr:
        cut = int((1 - fr / ir) / 2 * 100000)
        src.set("l", str(cut)); src.set("r", str(cut))
    else:
        cut = int((1 - ir / fr) / 2 * 100000)
        src.set("t", str(cut)); src.set("b", str(cut))
    return pic


def text(slide, ids, name, x, y, w, h, value, *, ea=GOTHIC, latin=LATIN, size=12.0, bold=False, color="tx1",
         line=None, spc=0, anchor="t", align="l", rot=0, wrap=True):
    """文言の欄。line は行間（pt）。spc は字間（pt）。"""
    tree = slide.shapes._spTree
    sp = sub(tree, "p:sp")
    nv = sub(sp, "p:nvSpPr")
    sub(nv, "p:cNvPr", id=ids.next(), name=name)
    sub(nv, "p:cNvSpPr", txBox="1")
    sub(nv, "p:nvPr")
    sppr = sub(sp, "p:spPr")
    xf = sub(sppr, "a:xfrm", **({"rot": int(rot * 60000)} if rot else {}))
    sub(xf, "a:off", x=E(x), y=E(y))
    sub(xf, "a:ext", cx=E(w), cy=E(h))
    g = sub(sppr, "a:prstGeom", prst="rect")
    sub(g, "a:avLst")
    sub(sppr, "a:noFill")
    tx = sub(sp, "p:txBody")
    bp = sub(tx, "a:bodyPr", wrap="square" if wrap else "none", lIns=0, tIns=0, rIns=0, bIns=0, anchor=anchor, rtlCol=0)
    sub(bp, "a:noAutofit")
    sub(tx, "a:lstStyle")
    for para in value.split("\n"):
        p = sub(tx, "a:p")
        ppr = sub(p, "a:pPr", algn=align, eaLnBrk=1, hangingPunct=1, latinLnBrk=0)
        if line:
            ls = sub(ppr, "a:lnSpc")
            sub(ls, "a:spcPts", val=int(line * 100))
        r = sub(p, "a:r")
        rpr = sub(r, "a:rPr", lang="ja-JP", altLang="en-US", sz=int(round(size * 100)), b=1 if bold else 0, dirty=0)
        if spc:
            rpr.set("spc", str(int(spc * 100)))
        _solid(rpr, color)
        sub(rpr, "a:latin", typeface=latin)
        sub(rpr, "a:ea", typeface=ea)
        t = sub(r, "a:t")
        t.text = para
    return sp


def notes(slide, yaml_text):
    slide.notes_slide.notes_text_frame.text = yaml_text


# ---------------------------------------------------------------- 型：表紙（PDF p1）
def cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    # 写真：左上の白い三角と、右の白い面に挟まれた五角形（PDF p1 の白い面2つの残り）
    picture(s, ids, "img.hero", ASSETS / "cover.jpg", 0, 0, 815, 539.5,
            pts=[(373, 0), (815, 0), (275.5, 539.5), (0, 539.5), (0, 373)], descr="表紙の写真")
    # 見出しの帯（黒 85%。PDF は黒の帯を写真に重ね、描画色は #262626 前後）
    poly(s, ids, "deco.headline-band", [(-8.5, 128.5), (513, 128.5), (428.5, 213), (-8.5, 213)], "tx1", alpha=85)
    poly(s, ids, "deco.subtitle-plate", [(438.5, 213), (523, 129), (872, 129), (872, 213)], "bg1")
    poly(s, ids, "deco.line-1", [(194.5, 158), (205, 158), (321.5, 41.5), (311, 41.5)], "accent1")
    poly(s, ids, "deco.line-2", [(520.5, 315.5), (531, 315.5), (596.5, 250), (586, 250)], "accent1")
    for i, x in enumerate([444, 606.5, 768.5], 1):
        rect(s, ids, f"deco.point-box-{i}", x, 326.5, 140.5, 140.5, "accent2")
        # 要点は PDF では 20pt だが、PDF はカナを詰めて組んである（palt）。PowerPoint は詰めないので
        # 「フルスクラッチ」が箱に入らない。箱の幅は変えられないので 17pt にする（D-015）
        text(s, ids, f"txt.point.{i}.lead", x + 10, 337, 125, 36, "ダミー小見出し", ea=GOTHIC_M, latin=LATIN_M, size=13.5, line=18)
        text(s, ids, f"txt.point.{i}.main", x + 10, 377, 125, 50, "ダミー要点", ea=GOTHIC_M, latin=LATIN_M, size=17, line=23)
    text(s, ids, "txt.headline", 95.5, 131, 345, 80, "ダミー大見出し\nダミー大見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, color="bg1")
    text(s, ids, "txt.subtitle", 525.5, 150, 432, 52, "ダミー副題\nダミー副題", ea=GOTHIC, latin=LATIN, size=15, bold=True, line=25.5)
    notes(s, "type: cover\nrole: [cover]\nlabel: 表紙（写真＋要点3つ）")


# ---------------------------------------------------------------- 型：課題提示（PDF p2）
def issues_frame(s, ids, eyebrow="Dummy Label"):
    poly(s, ids, "deco.corner", [(0, 0), (214, 0), (0, 214)], "accent1")
    poly(s, ids, "deco.corner-line", [(4.5, 219), (15, 219), (102.5, 131.5), (92, 131.5)], "accent1")
    poly(s, ids, "deco.title-band", [(108.5, 125.5), (192.5, 41), (953.5, 41), (953.5, 125.5)], "bg2")
    poly(s, ids, "deco.end-triangle", [(939.5, 449), (939.5, 522.5), (866, 522.5)], "accent2")
    # 斜めの英字（PDF は字間を空けた DIN を -45 度）
    text(s, ids, "txt.eyebrow", -32, 72, 240, 26, eyebrow, ea=GOTHIC, latin=LATIN, size=20.5, color="bg1",
         spc=4, line=26, rot=-45, anchor="ctr", align="ctr", wrap=False)
    text(s, ids, "txt.headline", 185.5, 52, 740, 62, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="ctr")


def issue_item(s, ids, i, x, y, img):
    """課題1件：左に写真（78pt角）、右に見出し・黄色の線・説明。"""
    picture(s, ids, f"img.item.{i}", img, x, y, 78, 78, descr=f"課題{i}の写真")
    # 見出しは2行まで（PDF p2「社長が現場から／離れられない」）。下揃えで、線の上に乗せる
    text(s, ids, f"txt.item.{i}.title", x + 90, y - 15, 274, 61, "ダミー課題", ea=GOTHIC_M, latin=LATIN_M, size=21, spc=3.5, line=29, anchor="b")
    # 情報の線（6-3「文字を直接強調する線は直線」）。平行四辺形ではなく細い長方形
    # 説明は PDF では 13pt（カナを詰めた太ゴ）。PowerPoint の幅で22字入るよう 12pt・幅274pt にし、線も合わせる（D-015）
    rect(s, ids, f"deco.item-{i}-rule", x + 90, y + 49, 274, 1.1, "accent2")
    text(s, ids, f"txt.item.{i}.body", x + 90, y + 58, 274, 20, "ダミー説明", ea=GOTHIC, latin=LATIN, size=12, bold=True, line=17)


def issues(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    issues_frame(s, ids)
    imgs = [ASSETS / f"item{k}.jpg" for k in (1, 2, 3, 4)]
    if n == 4:
        pos = [(160, 193), (554, 193), (160, 340), (554, 340)]
    elif n == 2:
        pos = [(160, 266), (554, 266)]
    else:  # 3件は1列に縦に並べる（2列だと右下が欠けて見えるため）
        pos = [(300, 160), (300, 265), (300, 370)]
    for i, (x, y) in enumerate(pos, 1):
        issue_item(s, ids, i, x, y, imgs[i - 1])
    notes(s, f'type: issues\nvariant: "{n}"\nrole: [problem]\nlabel: 課題提示（{n}項目）')


# ---------------------------------------------------------------- 型：数値訴求（PDF p3）
def numbers(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    picture(s, ids, "img.side", ASSETS / "side.jpg", 0, 0, 239, 540, descr="左の写真")
    # 写真に重ねる青（PDF は乗算。PowerPoint の図形は乗算を持たないので透明度で近づける。D-014）
    poly(s, ids, "deco.side-tri-1", [(0, 213.5), (239, -25.5), (239, 213.5)], "accent1", alpha=88)
    poly(s, ids, "deco.side-tri-2", [(0, 212.5), (239, 212.5), (0, 451.5)], "accent1", alpha=88)
    poly(s, ids, "deco.side-tri-3", [(78, 541.5), (239, 380.5), (239, 541.5)], "accent1", alpha=88)
    poly(s, ids, "deco.side-line", [(6, 223.5), (16.5, 223.5), (104, 136), (93.5, 136)], "accent2")
    poly(s, ids, "deco.top-band", [(744.5, 35.5), (818, -38), (1010, -38), (937, 35.5)], "accent2")
    poly(s, ids, "deco.slant-plate", [(379.5, 399.5), (646, 133), (879.5, 133), (613, 399.5)], "bg2", alpha=70)
    poly(s, ids, "deco.title-band", [(108.5, 125.5), (192.5, 41), (953.5, 41), (953.5, 125.5)], "bg2", alpha=92)
    poly(s, ids, "deco.end-triangle", [(939.5, 449), (939.5, 522.5), (866, 522.5)], "accent2")
    text(s, ids, "txt.eyebrow", -34, 58, 240, 56, "Dummy\nLabel", ea=GOTHIC, latin=LATIN, size=20.5, color="bg1",
         spc=4, line=26, rot=-45, anchor="ctr", align="ctr", wrap=False)
    text(s, ids, "txt.headline", 269, 52, 660, 62, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="ctr")
    # 数字の列の幅は PDF に合わせて真ん中を広く取る（2/3・2,000・70 の並び）。
    # DIN の代わりの Barlow は字幅が広いので、PDF の字の幅より少し広げてある
    for i, (x, vw) in enumerate([(296.5, 196), (497.5, 258), (762.5, 190)], 1):
        text(s, ids, f"txt.label.{i}", x, 180, 180, 32, "ダミー", ea=GOTHIC, latin=LATIN, size=23, bold=True, line=30)
        text(s, ids, f"num.value.{i}", x, 212, vw, 104, "00", ea=GOTHIC, latin=LATIN, size=99, line=104)
        poly(s, ids, f"deco.value-{i}-rule", [(x - 4.5, 310.5), (x + 141.5, 310.5), (x + 134.5, 303), (x + 3, 303)], "accent2")
        # 単位は黄色い線の右端にそろえる（PDF の「万円」の位置）
        text(s, ids, f"txt.unit.{i}", x, 312, 141.5, 28, "単位", ea=GOTHIC, latin=LATIN, size=23, bold=True, line=28, align="r")
        text(s, ids, f"txt.note.{i}", x + 6, 344, 180, 46, "ダミー説明", ea=GOTHIC, latin=LATIN, size=15, bold=True, line=22)
    text(s, ids, "txt.footnote", 300.5, 462, 600, 44, "ダミー注記", ea=GOTHIC, latin=LATIN, size=13, bold=True, line=22)
    notes(s, "type: numbers\nvariant: \"3\"\nrole: [number]\nlabel: 数値訴求（3つの数字）")


# ---------------------------------------------------------------- 型：注意・強調（PDF p5）
def attention(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    picture(s, ids, "img.hero", ASSETS / "cover.jpg", 0, 0, 656, 540,
            pts=[(418, 0), (656, 0), (116, 540), (0, 540), (0, 417.5)], descr="注意のページの写真")
    poly(s, ids, "deco.line-yellow", [(126.5, 263.5), (226, 164), (238.2, 164), (138.7, 263.5)], "accent2")
    poly(s, ids, "deco.line-blue", [(166.7, 514.8), (330, 351.5), (350, 351.5), (186.7, 514.8)], "accent1")
    poly(s, ids, "deco.band", [(290.3, 258.3), (374.6, 174), (973.8, 174), (973.8, 258.3)], "accent2")
    poly(s, ids, "deco.end-triangle", [(939.5, 449), (939.5, 522.5), (866, 522.5)], "accent2")
    text(s, ids, "txt.eyebrow", -28, 79, 240, 26, "Dummy Label", ea=GOTHIC, latin=LATIN, size=20.5, color="tx1",
         spc=4, line=26, rot=-45, anchor="ctr", align="ctr", wrap=False)
    # 見出しは黄色の帯の中に1行（PDF p5 の位置）
    text(s, ids, "txt.headline", 392, 186, 560, 60, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="ctr")
    # 説明は帯の下の白い面（PDF には無い。写真の斜めの辺より右に置く）
    text(s, ids, "txt.body", 412, 300, 500, 104, "ダミー説明", ea=GOTHIC_M, latin=LATIN_M, size=16, line=26)
    notes(s, "type: attention\nrole: [emphasis]\nlabel: 注意・強調（重要なことを改めて言う）")


# ---------------------------------------------------------------- 型：中扉（PDF p6）
def section(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    poly(s, ids, "deco.tri-left", [(0, 316.1), (92.7, 316.1), (0, 408.7)], "accent2")
    poly(s, ids, "deco.tri-right", [(963.7, 115), (963.7, 438.1), (640.6, 438.1)], "accent2")
    poly(s, ids, "deco.plate", [(85.3, 389.6), (302.6, 172.2), (833.4, 172.2), (615.1, 390.5)], "accent1")
    poly(s, ids, "deco.line", [(397.4, 540), (451.7, 485.7), (526.6, 485.7), (472.3, 540)], "accent1")
    text(s, ids, "txt.headline", 250, 251, 440, 60, "ダミー中扉", ea=GOTHIC, latin=LATIN, size=40, bold=True, line=50,
         color="bg1", anchor="ctr", align="ctr")
    notes(s, "type: section\nrole: [section]\nlabel: 中扉（章の見出し）")


# ---------------------------------------------------------------- 型：画像が並ぶ（手順）（PDF p7）
def steps(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    poly(s, ids, "deco.corner", [(0, 0), (430.7, 0), (0, 430.7)], "accent1")
    poly(s, ids, "deco.end-triangle", [(939.3, 449.2), (939.3, 522.7), (865.8, 522.7)], "accent2")
    text(s, ids, "txt.eyebrow", -34, 58, 240, 56, "Dummy Label", ea=GOTHIC_M, latin=LATIN, size=20.5, color="bg1",
         spc=4, line=26, rot=-45, anchor="ctr", align="ctr", wrap=False)
    # 見出しは PDF に無い。三角の右の白い面に置く
    text(s, ids, "txt.headline", 452, 64, 480, 56, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="ctr")
    imgs = [ASSETS / f"item{k}.jpg" for k in (1, 2, 3, 4)]
    if n == 4:
        size, xs, y = 163.6, [161.4, 349.2, 537.0, 724.8], 168.3
    else:
        size, xs, y = 200, [161.4, 424.9, 688.4], 150
    for i, x in enumerate(xs, 1):
        picture(s, ids, f"img.item.{i}", imgs[i - 1], x, y, size, size, descr=f"手順{i}の写真")
        rect(s, ids, f"deco.item-{i}-rule", x, y + size + 8, 28, 3, "accent1")
        text(s, ids, f"txt.item.{i}.caption", x, y + size + 18, size, 44, "ダミー説明", ea=GOTHIC, latin=LATIN, size=15, bold=True, line=21)
    notes(s, f'type: steps\nvariant: "{n}"\nrole: [steps]\nlabel: 画像が並ぶ（手順 {n}つ）')


# ---------------------------------------------------------------- 型：写真コラージュ（PDF p8）
DIAMONDS = [
    [(253.2, 169.3), (335.6, 86.8), (418.1, 169.3), (335.6, 251.8)],
    [(423.4, 169.4), (505.9, 87.0), (588.4, 169.4), (505.9, 251.9)],
    [(338.4, 254.5), (420.8, 172.0), (503.3, 254.5), (420.8, 337.0)],
    [(84.2, 339.5), (250.6, 173.1), (417.0, 339.5), (250.6, 506.0)],
    [(423.4, 339.5), (543.6, 219.3), (663.8, 339.5), (543.6, 459.8)],
    [(337.7, 425.2), (420.2, 342.7), (502.7, 425.2), (420.2, 507.7)],
]


def collage(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    poly(s, ids, "deco.band", [(546.6, 216.8), (763.4, 0), (960, 0), (960, 43.8), (666.8, 337)], "bg2")
    poly(s, ids, "deco.tri", [(716.4, 547.5), (969.7, 294.2), (969.7, 547.5)], "accent2")
    poly(s, ids, "deco.line-blue", [(632.4, 540), (708.3, 464.1), (726.1, 464.1), (650.2, 540)], "accent1")
    poly(s, ids, "deco.line-yellow", [(576.9, 136.2), (683.1, 30), (700.9, 30), (594.7, 136.2)], "accent2")
    picture(s, ids, "img.hero", ASSETS / "side.jpg", 84.2, 86.8, 579.6, 420.9, multi=DIAMONDS, descr="菱形のコラージュ")
    # 見出しは PDF では2行（6字ずつ）。菱形にかからないよう、上の白い面に1行で置く
    text(s, ids, "txt.headline", 48, 22, 520, 56, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="ctr")
    text(s, ids, "txt.subtitle", 50, 92, 196, 46, "ダミー副題", ea=GOTHIC, latin=LATIN, size=15, bold=True, line=21)
    notes(s, "type: collage\nrole: [closing]\nlabel: 写真コラージュ（結び・表紙の別案）")


# ---------------------------------------------------------------- 型：目次（PDF p9）
def toc(prs, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    poly(s, ids, "deco.corner", [(0, 0), (339.5, 0), (0, 339.5)], "accent1")
    poly(s, ids, "deco.tri", [(413.2, 545.7), (961.1, -2.2), (961.1, 545.7)], "accent2")
    poly(s, ids, "deco.line", [(752.5, 540), (899.3, 393.2), (919.3, 393.2), (772.5, 540)], "accent1")
    picture(s, ids, "img.side", ASSETS / "cover.jpg", 496.3, 208.2, 422.3, 218.3,
            pts=[(496.3, 426.5), (714.6, 208.2), (918.6, 208.2), (700.3, 426.5)], descr="目次の写真")
    text(s, ids, "txt.eyebrow", -28, 79, 240, 26, "Dummy Label", ea=GOTHIC, latin=LATIN, size=20.5, color="bg1",
         spc=4, line=26, rot=-45, anchor="ctr", align="ctr", wrap=False)
    # 項目は 45 度の流れに沿って左へずらす（青と黄の三角の間の白い面に収まる位置）
    y0 = 170 + (5 - n) * 31
    for k in range(n):
        y, x = y0 + k * 62, 230 - k * 40
        text(s, ids, f"deco.toc-no-{k + 1}", x, y, 52, 40, f"{k + 1:02d}", ea=GOTHIC, latin=LATIN_M, size=24, color="accent1", line=30, anchor="ctr")
        rect(s, ids, f"deco.toc-rule-{k + 1}", x + 60, y + 38, 300, 1.1, "accent2")
        text(s, ids, f"txt.item.{k + 1}", x + 60, y, 300, 36, "ダミー項目", ea=GOTHIC_M, latin=LATIN_M, size=18, line=24, anchor="ctr")
    notes(s, f'type: toc\nvariant: "{n}"\nrole: [toc]\nlabel: 目次（{n}項目）')


# ---------------------------------------------------------------- 型：写真＋本文A（PDF p10）
def photo_text_a(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    picture(s, ids, "img.right", ASSETS / "cover.jpg", 208.5, 0, 262.5, 540, descr="右の写真")
    picture(s, ids, "img.left", ASSETS / "item1.jpg", 31, 0, 177.5, 540, descr="左の写真")
    poly(s, ids, "deco.tri-top", [(208.5, 0), (470.4, 0), (208.5, 261.1)], "accent1", alpha=88)
    poly(s, ids, "deco.tri-bottom", [(75.5, 533.7), (208.5, 400.8), (208.5, 533.7)], "accent1", alpha=88)
    poly(s, ids, "deco.end-triangle", [(939.5, 449), (939.5, 522.5), (866, 522.5)], "accent2")
    text(s, ids, "txt.headline", 520, 70, 400, 80, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="b")
    rect(s, ids, "deco.rule", 520, 166, 60, 4, "accent2")
    text(s, ids, "txt.body", 520, 196, 400, 182, "ダミー本文", ea=GOTHIC, latin=LATIN, size=15, line=26)
    notes(s, 'type: photo-text\nvariant: "a"\nrole: [photo-text]\nlabel: 写真＋本文A（見出しと説明文）')


# ---------------------------------------------------------------- 型：写真＋本文B（PDF p11）
def photo_text_b(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    ids = Ids(s)
    picture(s, ids, "img.hero", ASSETS / "side.jpg", 0, 0, 506, 540, descr="左の写真")
    for name, pts in [("deco.tri-1", [(243.8, 6.2), (243.8, 268.3), (-18.3, 268.3)]),
                      ("deco.tri-2", [(-18.3, 529.2), (-18.3, 267.1), (243.8, 267.1)]),
                      ("deco.tri-3", [(506.4, 8.4), (506.4, 270.4), (244.3, 270.4)]),
                      ("deco.tri-4", [(244.3, 531.3), (244.3, 269.3), (506.4, 269.3)])]:
        poly(s, ids, name, pts, "accent1", alpha=88)
    text(s, ids, "txt.headline", 556, 70, 370, 80, "ダミー見出し", ea=MINCHO, latin=MINCHO, size=32.5, line=40, anchor="b")
    for k in range(3):
        y = 200 + k * 92
        rect(s, ids, f"deco.point-{k + 1}", 556, y + 5, 12, 12, "accent2")
        text(s, ids, f"txt.item.{k + 1}", 580, y, 346, 66, "ダミー要点", ea=GOTHIC_M, latin=LATIN_M, size=16, line=22)
    notes(s, 'type: photo-text\nvariant: "b"\nrole: [photo-text]\nlabel: 写真＋本文B（見出しと要点3つ）')


def main(out: Path):
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    set_theme(prs)
    cover(prs)
    for n in (2, 3, 4):
        issues(prs, n)
    numbers(prs)
    attention(prs)
    section(prs)
    for n in (3, 4):
        steps(prs, n)
    collage(prs)
    for n in (3, 4, 5):
        toc(prs, n)
    photo_text_a(prs)
    photo_text_b(prs)
    prs.save(str(out))
    print(out)


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "templates/provisional.pptx")
